from pathlib import Path
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook
import csv, re

def extract_text(path):
    p = Path(path)
    ext = p.suffix.lower()
    if ext == ".pdf":
        return "\n".join((page.extract_text() or "") for page in PdfReader(str(p)).pages), "pdf_text"
    if ext == ".docx":
        return "\n".join(x.text for x in Document(str(p)).paragraphs), "docx_text"
    if ext in {".txt", ".md"}:
        return p.read_text(encoding="utf-8", errors="ignore"), "plain_text"
    if ext == ".csv":
        with p.open("r", encoding="utf-8-sig", errors="ignore", newline="") as fh:
            return "\n".join(" | ".join(row) for row in csv.reader(fh)), "csv_text"
    if ext == ".xlsx":
        wb = load_workbook(str(p), read_only=True, data_only=True)
        lines = []
        for ws in wb.worksheets:
            lines.append(f"[SHEET: {ws.title}]")
            for row in ws.iter_rows(values_only=True):
                vals = [str(v) for v in row if v is not None]
                if vals:
                    lines.append(" | ".join(vals))
        return "\n".join(lines), "xlsx_text"
    if ext in {".png", ".jpg", ".jpeg", ".webp"}:
        try:
            from PIL import Image
            import pytesseract
            return pytesseract.image_to_string(Image.open(p)), "ocr"
        except Exception:
            return "", "ocr_unavailable"
    return "", "unsupported"

def norm(t):
    return " ".join((t or "").replace("\x00", " ").split())

def match(patterns, text):
    for pattern in patterns:
        m = re.search(pattern, text, re.I)
        if m:
            return m.group(1).strip()

def classify(text, filename):
    t = norm(text)
    l = t.lower()
    if any(x in l for x in ["insurance", "policy", "insured", "coverage", "certificate of insurance"]):
        typ = "insurance"
    elif any(x in l for x in ["building permit", "construction permit", "permit no", "permit number"]):
        typ = "permit"
    elif any(x in l for x in ["license", "licence", "contractor registration", "registration no"]):
        typ = "license"
    elif any(x in l for x in ["workers compensation", "workmen compensation", "worker compensation"]):
        typ = "workers_compensation"
    elif any(x in l for x in ["inspection certificate", "inspection report", "inspection authority"]):
        typ = "inspection_certificate"
    elif any(x in l for x in ["occupancy certificate", "completion certificate"]):
        typ = "completion_certificate"
    elif any(x in l for x in ["method statement", "safe work method statement", "swms"]):
        typ = "method_statement"
    elif any(x in l for x in ["risk assessment", "job hazard analysis", "jha"]):
        typ = "risk_assessment"
    elif any(x in l for x in ["environmental permit", "environmental clearance"]):
        typ = "environmental_permit"
    elif any(x in l for x in ["material certificate", "mill certificate", "test certificate"]):
        typ = "material_certificate"
    else:
        typ = "unknown"

    fields = {
        "policy_number": match([r"(?:policy|certificate)[\s#:_-]*(?:no\.?|number)?[\s#:_-]*([A-Z0-9/_-]{5,})"], t),
        "permit_number": match([r"(?:permit)[\s#:_-]*(?:no\.?|number)?[\s#:_-]*([A-Z0-9/_-]{4,})"], t),
        "license_number": match([r"(?:license|licence|registration)[\s#:_-]*(?:no\.?|number)?[\s#:_-]*([A-Z0-9/_-]{4,})"], t),
        "certificate_number": match([r"(?:certificate)[\s#:_-]*(?:no\.?|number)?[\s#:_-]*([A-Z0-9/_-]{4,})"], t),
        "expiry_date": match([
            r"(?:expiry|expires|expiration|valid\s+until|valid\s+through)[\s:,-]*([0-9]{1,2}[/-][0-9]{1,2}[/-][0-9]{2,4})",
            r"(?:expiry|expires|expiration|valid\s+until|valid\s+through)[\s:,-]*([A-Za-z]+\s+[0-9]{1,2},?\s+[0-9]{4})"
        ], t),
        "issue_date": match([
            r"(?:issue|issued|effective|commencement)[\s:,-]*(?:date)?[\s:,-]*([0-9]{1,2}[/-][0-9]{1,2}[/-][0-9]{2,4})"
        ], t),
        "named_insured": match([
            r"(?:named insured|insured|certificate holder)[\s:,-]+(.{3,100}?)(?:\s{2,}|policy|certificate|expiry|issue|coverage|limit)"
        ], t),
        "coverage_type": match([r"(?:coverage|policy type|insurance type)[\s:,-]+(.{3,80}?)(?:\s{2,}|limit|policy|expiry|issue)"], t),
        "coverage_limit": match([
            r"(?:limit|coverage limit|sum insured|insured amount)[\s:₹$€,-]*([0-9][0-9,]*(?:\.[0-9]+)?)"
        ], t),
        "authority": match([r"(?:authority|issuing authority|inspection authority)[\s:,-]+(.{3,100}?)(?:\s{2,}|project|issue|expiry|status)"], t),
    }
    return typ, fields, t
